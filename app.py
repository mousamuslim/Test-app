import os
import re
import shutil
import subprocess
import streamlit as st
from pptx import Presentation
from gTTS import gTTS
from pydub import AudioSegment
from PIL import Image, ImageDraw
from google import genai

st.set_page_config(page_title="PPTX to Video Generator", page_icon="🎬", layout="centered")

# جلب مفتاح Gemini من إعدادات Secrets بأمان
GEMINI_API_KEY = st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY"))
client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

def format_text_with_gemini(raw_text):
    prompt = f"""
    لديك النص التالي المستخرج من شريحة تعليمية:
    "{raw_text}"

    قم بإعادة صياغة النص وترتيبه بدقة بالنمط التالي فقط بدون أي زيادات أو مقدمات:
    الكلمة الأجنبية
    ترجمتها العربية
    الكلمة الأجنبية نفسها
    ترجمتها العربية مرة أخرى
    مثال على الكلمة باللغة الأجنبية
    ترجمة المثال للعربية
    المثال باللغة الأجنبية مرة أخرى
    ترجمة المثال للعربية مرة أخرى
    """
    response = client.models.generate_content(
        model='gemini-2.5-flash',
        contents=prompt,
    )
    return response.text.strip()

def generate_audio_for_formatted_text(formatted_text, output_audio_filename):
    lines = [line.strip() for line in formatted_text.split('\n') if line.strip()]
    combined_audio = AudioSegment.empty()
    
    for idx, line in enumerate(lines):
        is_arabic = bool(re.search(r'[\u0600-\u06FF]', line))
        lang = 'ar' if is_arabic else 'de'
        
        temp_file = f"temp_line_{idx}.mp3"
        tts = gTTS(text=line, lang=lang)
        tts.save(temp_file)
        
        segment = AudioSegment.from_mp3(temp_file)
        combined_audio += segment
        if os.path.exists(temp_file):
            os.remove(temp_file)
        
    combined_audio.export(output_audio_filename, format="mp3")
    return combined_audio

st.title("🎬 محول العروض التقديمية إلى فيديو وصوت")
st.write("قم برفع ملف PPTX لتحويله إلى فيديو مدمج ومقاطع صوتية تعليمية.")

uploaded_file = st.file_uploader("اختر ملف PowerPoint (.pptx)", type=["pptx"])

if uploaded_file and st.button("بدء المعالجة 🚀"):
    if not client:
        st.error("يرجى إدخال GEMINI_API_KEY في إعدادات Secrets الخاصة بـ Streamlit!")
    else:
        output_dir = "output"
        audio_dir = os.path.join(output_dir, "audio")
        video_dir = os.path.join(output_dir, "video_slides")
        img_dir = os.path.join(output_dir, "images")

        if os.path.exists(output_dir):
            shutil.rmtree(output_dir)

        os.makedirs(audio_dir, exist_ok=True)
        os.makedirs(video_dir, exist_ok=True)
        os.makedirs(img_dir, exist_ok=True)

        pptx_path = "input.pptx"
        with open(pptx_path, "wb") as f:
            f.write(uploaded_file.getbuffer())

        prs = Presentation(pptx_path)
        total_slides = len(prs.slides)
        
        progress_bar = st.progress(0)
        status_text = st.empty()
        
        full_presentation_audio = AudioSegment.empty()
        video_list = []

        for i, slide in enumerate(prs.slides):
            slide_num = i + 1
            status_text.text(f"جاري معالجة الشريحة {slide_num} من {total_slides}...")

            # 1. إنشاء صورة الشريحة
            img = Image.new('RGB', (1920, 1080), color=(245, 247, 250))
            draw = ImageDraw.Draw(img)
            text_lines = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            full_text = "\n".join(text_lines)
            draw.text((100, 100), f"Slide {slide_num}", fill=(30, 41, 59))
            draw.text((100, 250), full_text[:500], fill=(15, 23, 42))
            img_path = os.path.join(img_dir, f"slide_img_{slide_num}.png")
            img.save(img_path)

            # 2. التنسيق والصوت وفيديو الشريحة
            if full_text:
                formatted_text = format_text_with_gemini(full_text)
                audio_path = os.path.join(audio_dir, f"slide_{slide_num}.mp3")
                slide_audio = generate_audio_for_formatted_text(formatted_text, audio_path)
                full_presentation_audio += slide_audio

                video_path = os.path.join(video_dir, f"slide_video_{slide_num}.mp4")
                cmd = [
                    "ffmpeg", "-y",
                    "-loop", "1", "-i", img_path,
                    "-i", audio_path,
                    "-c:v", "libx264", "-tune", "stillimage",
                    "-c:a", "aac", "-b:a", "192k",
                    "-pix_fmt", "yuv420p", "-shortest",
                    video_path
                ]
                subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                video_list.append(video_path)

            progress_bar.progress((i + 1) / total_slides)

        status_text.text("جاري تجميع الملف النهائي...")

        # 3. حفظ الصوت المدمج
        full_audio_path = os.path.join(output_dir, "full_audio.mp3")
        full_presentation_audio.export(full_audio_path, format="mp3")

        # 4. حفظ الفيديو المدمج
        full_video_path = os.path.join(output_dir, "full_presentation.mp4")
        if video_list:
            files_txt = os.path.join(video_dir, "files.txt")
            with open(files_txt, "w") as f:
                for vid in video_list:
                    f.write(f"file '{os.path.basename(vid)}'\n")
            
            subprocess.run([
                "ffmpeg", "-y", "-f", "concat", "-safe", "0",
                "-i", "files.txt", "-c", "copy", "../full_presentation.mp4"
            ], cwd=video_dir, check=True)

        status_text.text("تمت المعالجة بنجاح! 🎉")
        st.success("تم تحويل العرض التقديمي بالكامل!")

        # 5. عرض النتائج وأزرار التنزيل
        st.subheader("📺 الفيديو الكامل")
        st.video(full_video_path)
        with open(full_video_path, "rb") as file:
            st.download_button("تنزيل الفيديو الكامل MP4", file, file_name="full_presentation.mp4", mime="video/mp4")

        st.subheader("🎧 الصوت الكامل")
        st.audio(full_audio_path)
        with open(full_audio_path, "rb") as file:
            st.download_button("تنزيل الصوت الكامل MP3", file, file_name="full_audio.mp3", mime="audio/mp3")
